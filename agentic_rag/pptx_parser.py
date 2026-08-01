"""Strukturierter PPTX-Parser — Locator v2.1.

Lesereihenfolge über rekursiven XY-Cut. Neu: Trennung von `content_text`
und `rendered_text` sowie versionierter Canonicalizer.
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from canonicalize import PPTX_SPEC, canonical_hash, canonicalize_units
from identity import (canonical_json, describe_document_version,
                      make_source_id)
from locators import slide_locator
from omml import M as MATH_NS, omml_to_latex

PARSER_NAME = "pptx-xycut"
PARSER_VERSION = "2.1"

GAP_MIN_X = 0.025
GAP_MIN_Y = 0.020


class Box:
    __slots__ = ("shape", "x0", "y0", "x1", "y1")

    def __init__(self, shape):
        self.shape = shape
        self.x0 = shape.left or 0
        self.y0 = shape.top or 0
        self.x1 = self.x0 + (shape.width or 0)
        self.y1 = self.y0 + (shape.height or 0)


def _widest_gap(boxes, axis: str):
    lo, hi = ("x0", "x1") if axis == "x" else ("y0", "y1")
    spans = sorted((getattr(b, lo), getattr(b, hi)) for b in boxes)
    merged = [list(spans[0])]
    for start, end in spans[1:]:
        if start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    if len(merged) < 2:
        return None
    return max(((merged[i + 1][0] - merged[i][1], merged[i][1])
                for i in range(len(merged) - 1)), key=lambda t: t[0])


def xy_cut(boxes, slide_w: int, slide_h: int, depth: int = 0) -> list:
    if len(boxes) <= 1 or depth > 6:
        return sorted(boxes, key=lambda b: (b.y0, b.x0))

    gx = _widest_gap(boxes, "x")
    gy = _widest_gap(boxes, "y")
    score_x = (gx[0] / slide_w) if gx else 0
    score_y = (gy[0] / slide_h) if gy else 0

    if score_x >= GAP_MIN_X and score_x >= score_y:
        cut = gx[1]
        left = [b for b in boxes if b.x0 <= cut]
        right = [b for b in boxes if b.x0 > cut]
        if left and right:
            return (xy_cut(left, slide_w, slide_h, depth + 1)
                    + xy_cut(right, slide_w, slide_h, depth + 1))

    if score_y >= GAP_MIN_Y:
        cut = gy[1]
        top = [b for b in boxes if b.y0 <= cut]
        bottom = [b for b in boxes if b.y0 > cut]
        if top and bottom:
            return (xy_cut(top, slide_w, slide_h, depth + 1)
                    + xy_cut(bottom, slide_w, slide_h, depth + 1))

    return sorted(boxes, key=lambda b: (b.y0, b.x0))


def _flatten(shapes) -> list:
    out = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(_flatten(sh.shapes))
        else:
            out.append(sh)
    return out


def _extract_math(shape, warnings: list) -> list:
    formulas = []
    try:
        element = shape._element
    except AttributeError:
        return formulas
    for node in element.iter():
        if isinstance(node.tag, str) and node.tag == MATH_NS + "oMath":
            latex, warn = omml_to_latex(node)
            if latex:
                formulas.append(latex)
            warnings.extend(warn)
    return formulas


def _shape_text(shape, warnings: list) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        alt = ""
        try:
            alt = shape._element._nvXxPr.cNvPr.get("descr", "") or ""
        except Exception:
            pass
        if not alt:
            warnings.append(f"Bild ohne Alternativtext (shape '{shape.name}') — "
                            f"möglicherweise eine als Grafik gesetzte Formel")
            return f"[BILD: {shape.name}, kein Alternativtext]"
        return f"[BILD: {alt}]"

    if shape.has_table:
        return "\n".join(" | ".join(c.text.strip() for c in row.cells)
                         for row in shape.table.rows)

    parts = []
    formulas = _extract_math(shape, warnings)
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    for latex in formulas:
        parts.append(f"$${latex}$$")
    return "\n".join(parts)


def parse_pptx(path: str, source_id: str = None):
    source_id = source_id or make_source_id("local", path)
    prs = Presentation(path)
    slide_w, slide_h = prs.slide_width, prs.slide_height

    docver = describe_document_version(
        path, source_id, PARSER_NAME, PARSER_VERSION,
        media_type="application/vnd.openxmlformats-officedocument."
                   "presentationml.presentation")
    dv_id = docver["document_version_id"]

    units = []
    for number, slide in enumerate(prs.slides, start=1):
        warnings: list = []
        shapes = _flatten(slide.shapes)

        title, body = [], []
        for sh in shapes:
            is_title = sh.is_placeholder and sh.placeholder_format.idx == 0
            (title if is_title else body).append(sh)

        boxes = [Box(sh) for sh in body
                 if sh.left is not None and sh.top is not None]
        undated = [sh for sh in body if sh.left is None or sh.top is None]
        if undated:
            warnings.append(f"{len(undated)} Shape(s) ohne Koordinaten — "
                            f"ans Ende gestellt")

        ordered = [b.shape for b in xy_cut(boxes, slide_w, slide_h)] + undated
        chunks = [t for t in (_shape_text(sh, warnings)
                              for sh in title + ordered) if t.strip()]

        content_text = "\n\n".join(chunks)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                content_text += f"\n\n[NOTIZEN]\n{notes}"

        loc = slide_locator(source_id, dv_id, number,
                            getattr(slide, "slide_id", None), content_text)
        loc.warnings = loc.warnings + warnings

        units.append({
            "locator": loc.to_dict(),
            "content_text": content_text,
            "rendered_text": f"{loc.header()}\n{content_text}",
            "n_shapes": len(chunks),
        })

    canonical_text, spec = canonicalize_units(
        [u["content_text"] for u in units], PPTX_SPEC)
    docver.update({
        "canonical_sha256": canonical_hash(canonical_text),
        "canonicalizer_name": spec["canonicalizer_name"],
        "canonicalizer_version": spec["canonicalizer_version"],
        "canonicalizer_spec_json": canonical_json(spec),
    })

    config = {"parser": PARSER_NAME, "parser_version": PARSER_VERSION,
              "reading_order": "xy-cut", "gap_min_x": GAP_MIN_X,
              "gap_min_y": GAP_MIN_Y, "include_notes": True,
              "omml_to_latex": True}
    return docver, units, config
